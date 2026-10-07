"""Read-only structured file inspection and explicit visual/OCR coverage."""
from pathlib import Path
import csv
import io
import shutil
import subprocess
import tempfile
import zipfile
import tarfile
from storage.file_store import get_file_record, resolve_file_path


def inspect_file(workspace_id, file_id, *, offset=0, limit=50):
    try:
        return _inspect_file(workspace_id, file_id, offset=offset, limit=limit)
    except (ValueError, OSError):
        raise
    except ImportError as exc:
        raise ValueError('document_processor_unavailable') from exc
    except Exception as exc:
        raise ValueError(f'document_inspection_failed:{type(exc).__name__}') from exc


def _inspect_file(workspace_id, file_id, *, offset=0, limit=50):
    if offset < 0 or not 1 <= limit <= 200:
        raise ValueError('invalid_inspection_range')
    record = get_file_record(workspace_id, file_id)
    if not record or record.get('lifecycle', 'active') != 'active':
        raise ValueError('file_unavailable')
    path = resolve_file_path(workspace_id, file_id)
    kind = record['file_kind']
    units, warnings, coverage = [], [], {}
    if kind == 'xlsx':
        from openpyxl import load_workbook
        formulas = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in formulas:
                values = cached[sheet.title]
                for row, cached_row in zip(sheet.iter_rows(), values.iter_rows()):
                    cells = [{'cell': cell.coordinate, 'value': value.value,
                              'formula': cell.value if cell.data_type == 'f' else None,
                              'source_value': cell.value} for cell, value in zip(row, cached_row) if cell.value is not None]
                    if cells:
                        units.append({'sheet': sheet.title, 'row': row[0].row, 'cells': cells})
            coverage = {'sheets': formulas.sheetnames, 'formulas': True, 'formula_calculation': False,
                        'charts_rendered': False, 'source_preserved': True}
            warnings.append('公式与已保存的缓存值分别展示。当前没有重算公式或渲染图表。')
        finally:
            formulas.close(); cached.close()
    elif kind == 'docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        for child in document.element.body.iterchildren():
            if child.tag.endswith('}p'):
                paragraph = Paragraph(child, document)
                units.append({'type': 'paragraph', 'text': paragraph.text, 'style': paragraph.style.name if paragraph.style else ''})
            elif child.tag.endswith('}tbl'):
                units.append({'type': 'table', 'rows': [[cell.text for cell in row.cells] for row in Table(child, document).rows]})
        with zipfile.ZipFile(path) as package:
            images = [n for n in package.namelist() if n.startswith('word/media/') and not n.endswith('/')]
        coverage = {'block_order': True, 'embedded_images': len(images), 'layout_rendered': False}
        warnings.append('结构预览保留正文与表格顺序；不还原 Word 分页和浮动对象版式。')
    elif kind == 'pptx':
        from pptx import Presentation
        presentation = Presentation(path)
        for page, slide in enumerate(presentation.slides, 1):
            shapes = []
            for shape in slide.shapes:
                entry = {'name': shape.name, 'type': str(shape.shape_type), 'bounds_emu': [shape.left, shape.top, shape.width, shape.height]}
                if shape.has_text_frame:
                    entry['text'] = shape.text
                if shape.has_table:
                    entry['rows'] = [[cell.text for cell in row.cells] for row in shape.table.rows]
                shapes.append(entry)
            units.append({'page': page, 'shapes': shapes,
                          'notes': slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''})
        coverage = {'slides': len(presentation.slides), 'shape_structure': True, 'layout_rendered': False,
                    'size_emu': [presentation.slide_width, presentation.slide_height]}
        warnings.append('已读取幻灯片文字、表格、备注和对象位置；尚未渲染完整幻灯片版式。')
    elif kind == 'pdf':
        import pdfplumber
        with pdfplumber.open(path) as document:
            units = [{'page': n, 'text': page.extract_text() or '', 'width': page.width, 'height': page.height} for n, page in enumerate(document.pages, 1)]
        coverage = {'pages': len(units), 'text_pages': sum(bool(p['text'].strip()) for p in units),
                    'page_render': True, 'ocr_available': bool(shutil.which('tesseract')), 'ocr_applied': False}
        if coverage['text_pages'] < len(units):
            warnings.append('部分页面没有提取到文字。可以读取页图，或在处理器可用时明确执行 OCR。')
    elif kind == 'csv':
        text = path.read_text(encoding='utf-8-sig', errors='replace')
        try:
            dialect = csv.Sniffer().sniff(text[:8192])
        except csv.Error:
            dialect = csv.excel_tab if path.suffix.lower() == '.tsv' else csv.excel
        units = [{'row': n, 'cells': row} for n, row in enumerate(csv.reader(io.StringIO(text), dialect), 1)]
        coverage = {'rows': len(units), 'delimiter': dialect.delimiter}
    elif kind == 'zip':
        with zipfile.ZipFile(path) as archive:
            units = [{'name': m.filename, 'size_bytes': m.file_size, 'compressed_bytes': m.compress_size, 'directory': m.is_dir()} for m in archive.infolist()]
        coverage = {'members': len(units), 'extracted': False}
    elif kind in {'tar', 'gz', 'bz2'} and tarfile.is_tarfile(path):
        with tarfile.open(path) as archive:
            units = [{'name': m.name, 'size_bytes': m.size, 'directory': m.isdir(), 'type': 'file' if m.isfile() else 'directory' if m.isdir() else 'link_or_special'} for m in archive]
        coverage = {'members': len(units), 'extracted': False}
    elif not record.get('binary'):
        units = [{'line': n, 'text': line} for n, line in enumerate(path.read_text(encoding='utf-8', errors='replace').splitlines(), 1)]
        coverage = {'lines': len(units)}
    else:
        raise ValueError('no_structured_processor_for_file_kind')
    end = min(len(units), offset + limit)
    return {'file_id': file_id, 'file_kind': kind, 'units': units[offset:end], 'offset': offset,
            'next_offset': end if end < len(units) else None, 'total_units': len(units),
            'coverage': coverage, 'warnings': warnings, 'source_sha256': record['sha256']}


def render_pdf_page(workspace_id, file_id, *, page=1, ocr=False):
    record = get_file_record(workspace_id, file_id)
    if not record or record.get('lifecycle', 'active') != 'active' or record['file_kind'] != 'pdf':
        raise ValueError('active_pdf_required')
    if ocr and not shutil.which('tesseract'):
        raise ValueError('ocr_processor_unavailable')
    import pypdfium2
    with pypdfium2.PdfDocument(resolve_file_path(workspace_id, file_id)) as document:
        if not 1 <= page <= len(document):
            raise ValueError('invalid_pdf_page')
        source = document[page - 1]
        try:
            width, height = source.get_size()
            bitmap = source.render(scale=min(2, 2400 / max(width, height)))
            try:
                image = bitmap.to_pil()
                output = io.BytesIO(); image.save(output, format='PNG'); image.close()
            finally:
                bitmap.close()
        finally:
            source.close()
        count = len(document)
    from storage.file_store import import_user_upload
    image_record = import_user_upload(workspace_id, io.BytesIO(output.getvalue()), f'{Path(record["original_name"]).stem}_page_{page}.png',
        source='document_page_render', logical_type='tmp', metadata={'source_file_id': file_id, 'page': page, 'source_sha256': record['sha256']})
    from storage.reference_index import add_reference
    add_reference(workspace_id, image_record.file_id, 'file', file_id, 'page_image')
    text = ''
    if ocr:
        with tempfile.TemporaryDirectory(prefix='lzcore-ocr-') as directory:
            target = Path(directory) / 'page.png'; target.write_bytes(output.getvalue())
            result = subprocess.run([shutil.which('tesseract'), str(target), 'stdout'], capture_output=True, text=True, timeout=60)
            if result.returncode:
                raise ValueError('ocr_processor_failed')
            text = result.stdout
    return {'file_id': file_id, 'page': page, 'page_count': count, 'image_file_id': image_record.file_id,
            'ocr_applied': ocr, 'text': text, 'source_sha256': record['sha256']}
