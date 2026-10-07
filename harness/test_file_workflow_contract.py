"""File workflows exercise actual payloads, owner persistence and recovery."""
import io
import pytest


@pytest.fixture
def files_ws(tmp_path, monkeypatch):
    monkeypatch.setenv('LZCORE_WORKSPACE_ROOT', str(tmp_path))
    monkeypatch.setenv('LZCORE_WORKSPACE_DIR', str(tmp_path))
    return 'file_contract'


@pytest.mark.parametrize('name,kind,binary', [
    ('数据.csv', 'csv', False), ('report.xlsx', 'xlsx', True),
    ('图.svg', 'svg', True), ('music.mp3', 'mp3', True),
    ('bundle.zip', 'zip', True), ('unknown.custom', 'binary', True),
])
def test_upload_and_tool_import_share_classification(files_ws, name, kind, binary):
    from storage.file_store import import_user_upload, resolve_file_path
    from backend.api.artifact_routes import _guess_upload_kind
    payload = b'opaque\x00bytes'
    record = import_user_upload(files_ws, io.BytesIO(payload), name)
    assert (record.file_kind, record.binary) == (kind, binary) == _guess_upload_kind(name)
    assert resolve_file_path(files_ws, record.file_id).read_bytes() == payload


def test_message_attachment_replacement_and_session_deletion(files_ws):
    from storage.file_store import import_user_upload
    from storage.message_store import SessionMessageStore
    from storage.reference_index import list_references_for_file
    from storage.session_store import delete_session_permanently
    a = import_user_upload(files_ws, io.BytesIO(b'a'), 'a.txt')
    b = import_user_upload(files_ws, io.BytesIO(b'b'), 'b.txt')
    first = SessionMessageStore('session_first', files_ws)
    other = SessionMessageStore('session_other', files_ws)
    for store in (first, first, other):
        store.write_message('run_a', 'user', 'use file', metadata={'attachments': [{'file_id': a.file_id}]})
    assert len(list_references_for_file(files_ws, a.file_id)) == 2
    first.write_message('run_a', 'user', 'replace file', metadata={'attachments': [{'file_id': b.file_id}]})
    assert len(list_references_for_file(files_ws, a.file_id)) == 1
    assert len(list_references_for_file(files_ws, b.file_id)) == 1
    assert delete_session_permanently('session_first', files_ws, confirm=True)
    assert not list_references_for_file(files_ws, b.file_id)
    assert len(list_references_for_file(files_ws, a.file_id)) == 1


def test_reference_repeated_save_is_idempotent(files_ws):
    from storage.reference_index import add_reference, list_references_for_file
    first = add_reference(files_ws, 'file_example', 'artifact', 'art_example', 'content')
    second = add_reference(files_ws, 'file_example', 'artifact', 'art_example', 'content', {'used': True})
    assert first.ref_id == second.ref_id
    assert list_references_for_file(files_ws, 'file_example')[0]['metadata']['used']


def test_mutation_keeps_referenced_bytes_and_publishes_working_version(files_ws):
    from storage.file_store import import_user_upload, resolve_file_path, get_file_record
    from storage.reference_index import add_reference
    from storage.file_mutations import managed_file_mutation
    from storage.project_changes import workspace_files_lock
    record = import_user_upload(files_ws, io.BytesIO(b'original'), 'source.txt')
    path = resolve_file_path(files_ws, record.file_id)
    add_reference(files_ws, record.file_id, 'message', 'session/run:user', 'attachment')
    with workspace_files_lock(files_ws), managed_file_mutation(files_ws, paths=[path]) as changes:
        path.write_bytes(b'updated work')
    assert resolve_file_path(files_ws, record.file_id).read_bytes() == b'original'
    assert changes[0]['file_id'] != record.file_id
    current = get_file_record(files_ws, changes[0]['file_id'])
    assert current['size_bytes'] == 12
    assert current['metadata']['previous_file_id'] == record.file_id
    assert resolve_file_path(files_ws, current['file_id']).read_bytes() == b'updated work'


def test_noop_mutation_does_not_create_another_file(files_ws):
    from storage.file_store import import_user_upload, resolve_file_path, list_files
    from storage.reference_index import add_reference
    from storage.file_mutations import managed_file_mutation
    record = import_user_upload(files_ws, io.BytesIO(b'original'), 'source.txt')
    add_reference(files_ws, record.file_id, 'message', 'session/run:user')
    with managed_file_mutation(files_ws, paths=[resolve_file_path(files_ws, record.file_id)]) as changes:
        pass
    assert changes == []
    assert len(list_files(files_ws)) == 1


def test_recycle_restore_and_permanent_clear_preserve_owner_identity(files_ws):
    from storage.file_store import import_user_upload, get_file_record, restore_file, resolve_file_path
    from storage.reference_index import add_reference, list_references_for_file
    from storage.data_management import delete_unreferenced_file
    record = import_user_upload(files_ws, io.BytesIO(b'original'), 'source.txt')
    add_reference(files_ws, record.file_id, 'message', 'session/run:user', 'attachment')
    assert not delete_unreferenced_file(files_ws, record.file_id)['ok']
    recycled = delete_unreferenced_file(files_ws, record.file_id, force=True)
    assert recycled['recoverable'] and recycled['lifecycle'] == 'soft_deleted'
    assert resolve_file_path(files_ws, record.file_id).read_bytes() == b'original'
    assert restore_file(files_ws, record.file_id)['ok']
    assert restore_file(files_ws, record.file_id)['already_active']
    purged = delete_unreferenced_file(files_ws, record.file_id, force=True, permanent=True)
    assert not purged['recoverable']
    assert get_file_record(files_ws, record.file_id)['lifecycle'] == 'purged'
    assert len(list_references_for_file(files_ws, record.file_id)) == 1
    assert not restore_file(files_ws, record.file_id)['ok']


def test_preview_range_can_continue_to_full_content(files_ws):
    from storage.file_store import import_user_upload
    from storage.data_management import text_file_content
    record = import_user_upload(files_ws, io.BytesIO('中文abcdef'.encode()), 'text.txt')
    first = text_file_content(files_ws, record.file_id, max_chars=3)
    assert first['truncated'] and first['next_offset'] == 3
    second = text_file_content(files_ws, record.file_id, max_chars=10, offset=first['next_offset'])
    assert first['content'] + second['content'] == '中文abcdef'
    assert not second['truncated']


@pytest.fixture
def file_client(files_ws):
    from core.tools.client import ToolRuntimeClient
    from core.tools.registry import ToolRegistry
    from core.tools.canonical_registry import to_tool_specs
    registry = ToolRegistry()
    for spec, handler in to_tool_specs():
        if spec.tool_id in {'workspace.filestore', 'workspace.file'}:
            registry.register_tool(spec, handler)
    return ToolRuntimeClient(registry)


def test_governed_materialize_and_binary_publish(file_client, files_ws):
    from core.tools.context import ToolRuntimeContext
    from storage.file_store import import_user_upload, resolve_file_path
    from storage.reference_index import list_references_for_file
    import openpyxl
    context = ToolRuntimeContext(workspace_id=files_ws, requested_by='turn_runner', run_id='run_file', session_id='session_file')
    original = import_user_upload(files_ws, io.BytesIO(b'input'), 'input.txt')
    def invoke(action, **args):
        result = file_client.invoke('workspace.filestore', {'action': action, **args}, context=context)
        assert result.status == 'succeeded', result.errors
        return result.output
    material = invoke('materialize', file_id=original.file_id, destination='files/data/project/input.txt')
    assert material['filepath'] == 'files/data/project/input.txt'
    from storage.paths import workspace_root
    book_path = workspace_root(files_ws) / 'files/data/project/result.xlsx'
    book = openpyxl.Workbook(); book.active['A1'] = '=1+2'; book.save(book_path)
    published = invoke('publish', filepath='files/data/project/result.xlsx', source_file_ids=[original.file_id])
    assert resolve_file_path(files_ws, published['file_id']).read_bytes() == book_path.read_bytes()
    assert openpyxl.load_workbook(resolve_file_path(files_ws, published['file_id'])).active['A1'].value == '=1+2'
    assert any(r['owner_id'] == published['artifact_id'] and r['relation'] == 'source' for r in list_references_for_file(files_ws, original.file_id))
    assert resolve_file_path(files_ws, original.file_id).read_bytes() == b'input'
    repeat = file_client.invoke('workspace.filestore', {'action': 'materialize', 'file_id': original.file_id, 'destination': material['filepath']}, context=context)
    assert repeat.status != 'succeeded'
    assert (workspace_root(files_ws) / material['filepath']).read_bytes() == b'input'


def test_files_cursor_and_organization_preserve_identity(files_ws):
    from storage.file_store import import_user_upload
    from storage.file_workspace import files_page, organize_file
    records = [import_user_upload(files_ws, io.BytesIO(str(n).encode()), f'{n}.txt') for n in range(15)]
    seen, cursor = [], ''
    while True:
        page = files_page(files_ws, limit=4, cursor=cursor)
        seen += [r['file_id'] for r in page['files']]
        cursor = page['next_cursor']
        if not cursor: break
    assert len(seen) == len(set(seen)) == 15
    organized = organize_file(files_ws, records[0].file_id, name='新名称.txt', folder='项目/资料')
    assert organized['file_id'] == records[0].file_id
    assert files_page(files_ws, folder='项目/资料')['files'][0]['file_id'] == records[0].file_id
    with pytest.raises(ValueError): organize_file(files_ws, records[0].file_id, folder='../escape')


def test_read_image_produces_real_provider_evidence(file_client, files_ws):
    from PIL import Image
    from storage.paths import workspace_root, ensure_workspace_storage_dirs
    from core.tools.context import ToolRuntimeContext
    from agent.runtime.vision_inputs import build_vision_content
    from core.runtime_engine.evidence import evidence_to_vision_references
    ensure_workspace_storage_dirs(files_ws)
    image = workspace_root(files_ws) / 'files/data/diagram.png'
    Image.new('RGB', (8, 8), color='navy').save(image)
    result = file_client.invoke('workspace.file', {'action': 'read_image', 'filepath': 'files/data/diagram.png'},
        context=ToolRuntimeContext(workspace_id=files_ws, requested_by='turn_runner'))
    assert result.status == 'succeeded'
    parts, warnings = build_vision_content(evidence_to_vision_references(result.output['evidence_parts']), files_ws)
    assert not warnings
    assert parts[0]['image_url']['url'].startswith('data:image/png;base64,')


def test_format_inspection_preserves_formulas_structure_and_ranges(files_ws):
    from openpyxl import Workbook
    from docx import Document
    from pptx import Presentation
    from storage.file_store import import_user_upload
    from storage.file_formats import inspect_file
    book = Workbook(); book.active['A1'] = 2; book.active['B1'] = '=A1*3'
    raw = io.BytesIO(); book.save(raw)
    xlsx = import_user_upload(files_ws, io.BytesIO(raw.getvalue()), '公式.xlsx')
    result = inspect_file(files_ws, xlsx.file_id)
    assert result['units'][0]['cells'][1]['formula'] == '=A1*3'
    assert result['units'][0]['cells'][1]['value'] is None
    assert not result['coverage']['formula_calculation']
    document = Document(); document.add_paragraph('before'); document.add_table(1, 1).cell(0, 0).text = 'middle'; document.add_paragraph('after')
    raw = io.BytesIO(); document.save(raw)
    docx = import_user_upload(files_ws, io.BytesIO(raw.getvalue()), '结构.docx')
    result = inspect_file(files_ws, docx.file_id, limit=2)
    assert [u['type'] for u in result['units']] == ['paragraph', 'table']
    assert inspect_file(files_ws, docx.file_id, offset=result['next_offset'])['units'][0]['text'] == 'after'
    slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[1]); slide.shapes.title.text = 'Title'
    raw = io.BytesIO(); slides.save(raw)
    pptx = import_user_upload(files_ws, io.BytesIO(raw.getvalue()), '页.pptx')
    assert inspect_file(files_ws, pptx.file_id)['units'][0]['shapes'][0]['text'] == 'Title'


def test_pdf_page_render_actual_bytes_and_invalid_page(files_ws):
    from PIL import Image
    from storage.file_store import import_user_upload, resolve_file_path
    from storage.file_formats import render_pdf_page, inspect_file
    raw = io.BytesIO(); Image.new('RGB', (100, 80), 'red').save(raw, format='PDF')
    pdf = import_user_upload(files_ws, io.BytesIO(raw.getvalue()), '扫描.pdf')
    assert inspect_file(files_ws, pdf.file_id)['coverage']['text_pages'] == 0
    page = render_pdf_page(files_ws, pdf.file_id)
    with Image.open(resolve_file_path(files_ws, page['image_file_id'])) as image:
        assert image.format == 'PNG' and image.width >= 100
    with pytest.raises(ValueError, match='invalid_pdf_page'):
        render_pdf_page(files_ws, pdf.file_id, page=2)


def test_source_move_and_archive_roundtrip_preserve_bytes(files_ws):
    from storage.workspace_files import create_workspace_text
    from storage.source_workspace import source_entries, move_source, pack_source, extract_archive
    from storage.file_store import import_user_upload, get_file_record
    from storage.paths import workspace_root
    create_workspace_text(files_ws, 'files/data/project/中文 source.txt', 'actual source')
    entries = source_entries(files_ws, 'files/data/project')['entries']
    assert entries[0]['filepath'] == 'files/data/project/中文 source.txt'
    move_source(files_ws, entries[0]['filepath'], 'files/data/project/moved.txt')
    assert not (workspace_root(files_ws) / entries[0]['filepath']).exists()
    result = pack_source(files_ws, 'files/data/project', 'files/data/package.zip')
    record = import_user_upload(files_ws, workspace_root(files_ws) / result['filepath'], 'package.zip')
    extract_archive(files_ws, record.file_id, 'files/data/extracted')
    assert (workspace_root(files_ws) / 'files/data/extracted/moved.txt').read_text() == 'actual source'
    with pytest.raises(FileExistsError):
        extract_archive(files_ws, record.file_id, 'files/data/extracted')


@pytest.mark.parametrize('member', ['../escape.txt', '/absolute.txt', 'C:/escape.txt', 'a\\b.txt', 'nul\x00suffix.txt'])
def test_archive_traversal_rejected_before_writing(files_ws, member):
    import zipfile
    from storage.file_store import import_user_upload
    from storage.source_workspace import extract_archive
    from storage.paths import workspace_root
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, 'w') as archive:
        info = zipfile.ZipInfo('fixture')
        info.filename = member  # Keep adversarial bytes intact on Windows too.
        archive.writestr(info, 'no')
    record = import_user_upload(files_ws, io.BytesIO(raw.getvalue()), 'bad.zip')
    with pytest.raises(ValueError, match='unsafe_archive_member'):
        extract_archive(files_ws, record.file_id, 'files/data/destination')
    assert not (workspace_root(files_ws) / 'files/data/destination').exists()


def test_non_native_image_transport_reuses_actual_png(files_ws):
    from PIL import Image
    from storage.paths import workspace_root
    from storage.image_transport import image_for_evidence
    from storage.file_store import resolve_file_path
    path = workspace_root(files_ws) / 'files/data/image.bmp'; path.parent.mkdir(parents=True, exist_ok=True)
    Image.new('RGB', (10, 12), 'blue').save(path)
    first, dimensions = image_for_evidence(files_ws, path)
    assert image_for_evidence(files_ws, path)[0] == first
    assert dimensions == [10, 12]
    with Image.open(resolve_file_path(files_ws, first)) as image:
        assert image.format == 'PNG'


def test_content_search_versions_scope_and_literal_query(files_ws):
    from storage.file_store import import_user_upload
    from storage.file_search import search_content, synchronize
    from storage.principal import storage_principal
    record = import_user_upload(files_ws, io.BytesIO('第一行\n正文找到 100%_literal\n第三行'.encode()), '内容.txt')
    hit = search_content(files_ws, '100%_literal')['hits'][0]
    assert hit['file_id'] == record.file_id and hit['line'] == 2
    assert search_content(files_ws, '不存在')['total'] == 0
    assert not synchronize(files_ws)['updated']
    with storage_principal('another-user'):
        assert search_content(files_ws, '100%_literal')['total'] == 0


def test_reference_migration_preview_apply_is_idempotent(files_ws):
    from storage.file_store import import_user_upload
    from storage.message_store import SessionMessageStore
    from storage.reference_index import list_references
    from storage.file_audit import migrate_references, file_health
    from storage.paths import workspace_root
    record = import_user_upload(files_ws, io.BytesIO(b'actual'), 'source.txt')
    SessionMessageStore('session_source', files_ws).write_message('run_input', 'user', 'use', metadata={'attachments': [{'file_id': record.file_id}]})
    (workspace_root(files_ws) / 'index/references.jsonl').write_text('')
    report = migrate_references(files_ws)
    assert len(report['additions']) == 1 and not list_references(files_ws)
    applied = migrate_references(files_ws, apply=True)
    assert applied['backup'] and len(list_references(files_ws)) == 1
    assert not migrate_references(files_ws, apply=True)['additions']
    assert file_health(files_ws, hashes=True)['ok']


def test_file_bundle_roundtrip_conflict_and_idempotent_restore(files_ws):
    import shutil
    from storage.file_store import import_user_upload, resolve_file_path
    from storage.message_store import SessionMessageStore
    from storage.file_bundle import export_bundle, restore_bundle
    from storage.paths import workspace_root
    from storage.reference_index import list_references
    record = import_user_upload(files_ws, io.BytesIO(b'\x00original office bytes'), 'opaque.custom')
    SessionMessageStore('session_bundle', files_ws).write_message('run_input', 'user', 'use original', metadata={'attachments': [{'file_id': record.file_id}]})
    bundle = export_bundle(files_ws)
    shutil.rmtree(workspace_root(files_ws))
    assert restore_bundle(files_ws, bundle)['preview']
    assert restore_bundle(files_ws, bundle, preview=False)['restored']
    assert resolve_file_path(files_ws, record.file_id).read_bytes() == b'\x00original office bytes'
    assert len(list_references(files_ws)) == 1
    assert restore_bundle(files_ws, bundle, preview=False)['restored']
    resolve_file_path(files_ws, record.file_id).write_bytes(b'conflict')
    assert not restore_bundle(files_ws, bundle, preview=False)['ok']
    assert resolve_file_path(files_ws, record.file_id).read_bytes() == b'conflict'
    with pytest.raises(ValueError, match='backup_scope'):
        restore_bundle('another-workspace', bundle)


def test_unknown_file_commit_reconciles_bytes_without_source_replay(files_ws, monkeypatch):
    from storage import index
    from storage.file_store import import_user_upload, FileCommitUnknown, get_file_record
    from storage.file_audit import reconcile_file_commits
    original = index.append_file_record
    def fail(*args):
        raise RuntimeError('injected metadata failure')
    monkeypatch.setattr(index, 'append_file_record', fail)
    with pytest.raises(FileCommitUnknown) as error:
        import_user_upload(files_ws, io.BytesIO(b'one original write'), '中断.txt')
    fid = error.value.record['file_id']
    assert error.value.as_result()['automatic_retry_allowed'] is False
    assert not get_file_record(files_ws, fid)
    assert reconcile_file_commits(files_ws)['results'][0]['state'] == 'verified_payload_unindexed'
    monkeypatch.setattr(index, 'append_file_record', original)
    result = reconcile_file_commits(files_ws, apply=True)
    assert result['payload_writes'] == 0
    assert get_file_record(files_ws, fid)['sha256'] == error.value.record['sha256']
    assert not reconcile_file_commits(files_ws, apply=True)['results']


def test_governed_file_write_unknown_retains_recovery_identity(file_client, files_ws, monkeypatch):
    from storage import index
    from core.tools.context import ToolRuntimeContext
    monkeypatch.setattr(index, 'append_file_record', lambda *_: (_ for _ in ()).throw(RuntimeError('fault after payload')))
    result = file_client.invoke('workspace.file', {'action': 'write', 'filename': 'out.txt', 'content': 'real bytes'}, context=ToolRuntimeContext(workspace_id=files_ws, requested_by='turn_runner'))
    assert result.status == 'failed'
    assert result.output['error_code'] == 'EXECUTION_UNKNOWN'
    assert result.output['file_id'].startswith('file_')
    assert not result.output['automatic_retry_allowed']


def test_storage_http_auth_payload_and_workspace_isolation(files_ws, monkeypatch):
    from backend.main import create_app
    from storage.principal import storage_principal
    from storage.file_store import import_user_upload
    monkeypatch.setenv('LZCORE_AUTH_ENABLED', 'true')
    monkeypatch.setenv('LZCORE_API_TOKEN', 'file-workflow-fixture-token')
    monkeypatch.setenv('LZCORE_LOGIN_ENABLED', 'false')
    monkeypatch.setenv('LZCORE_IDENTITY_ENABLED', 'false')
    app = create_app(); client = app.test_client()
    headers = {'Authorization': 'Bearer file-workflow-fixture-token'}
    assert client.get('/api/storage/files', query_string={'workspace_id': files_ws}).status_code == 401
    upload = client.post(f'/api/workspaces/{files_ws}/artifacts/upload', data={'file': (io.BytesIO(b'\x00opaque bytes'), '新类型.custom')}, headers=headers).get_json()
    fid = upload['file']['file_id']
    assert client.get(f'/api/storage/files/{fid}/download', query_string={'workspace_id': files_ws}, headers=headers).data == b'\x00opaque bytes'
    assert client.get(f'/api/storage/files/{fid}/download', query_string={'workspace_id': 'other_workspace'}, headers=headers).status_code == 404
    with storage_principal('another-user'):
        private = import_user_upload(files_ws, io.BytesIO(b'private'), 'private.txt')
    assert client.get(f'/api/storage/files/{private.file_id}/download', query_string={'workspace_id': files_ws}, headers=headers).status_code == 404
    renamed = client.patch(f'/api/storage/files/{fid}', json={'workspace_id': files_ws, 'name': '重命名.custom', 'folder': '项目/资料'}, headers=headers)
    assert renamed.status_code == 200 and renamed.get_json()['file']['file_id'] == fid
    deleted = client.delete(f'/api/storage/files/{fid}', query_string={'workspace_id': files_ws, 'confirm': 'true'}, headers=headers)
    assert deleted.get_json()['recoverable']
    assert client.post(f'/api/storage/files/{fid}/restore', json={'workspace_id': files_ws}, headers=headers).get_json()['ok']


def test_artifact_commit_failure_reconciles_original_identity_and_references(files_ws, monkeypatch):
    from storage.file_store import import_user_upload, FileCommitUnknown, resolve_file_path
    from storage import artifact_metadata_store as metadata
    from storage.file_audit import reconcile_file_commits, file_health
    from storage.reference_index import list_references_for_owner
    original = import_user_upload(files_ws, io.BytesIO(b'original payload'), '原件.bin')
    append = metadata._append_index_unlocked
    monkeypatch.setattr(metadata, '_append_index_unlocked', lambda *_: (_ for _ in ()).throw(RuntimeError('projection fault')))
    with pytest.raises(FileCommitUnknown) as error:
        metadata.create_artifact_metadata(workspace_id=files_ws, file_record=original,
            artifact_type='output_data', title='原件交付', metadata={'source_file_ids': [original.file_id]})
    aid = error.value.as_result()['artifact_id']
    assert aid.startswith('art_')
    assert not list_references_for_owner(files_ws, 'artifact', aid)
    assert reconcile_file_commits(files_ws)['results'][0]['state'] == 'verified_artifact_commit'
    monkeypatch.setattr(metadata, '_append_index_unlocked', append)
    result = reconcile_file_commits(files_ws, apply=True)
    assert result['payload_writes'] == 0
    assert resolve_file_path(files_ws, original.file_id).read_bytes() == b'original payload'
    assert {r['relation'] for r in list_references_for_owner(files_ws, 'artifact', aid)} == {'content', 'source'}
    assert file_health(files_ws, hashes=True)['ok']
    assert not reconcile_file_commits(files_ws, apply=True)['results']


def test_knowledge_replacement_preserves_referenced_original_bytes(files_ws):
    from storage.file_store import write_knowledge_document, read_file_content, get_file_record
    from storage.reference_index import add_reference
    source = 'ksrc_0123456789ab'
    first = write_knowledge_document(files_ws, source, 'old knowledge', title='说明')
    add_reference(files_ws, first.file_id, 'message', 's/r:user', 'attachment')
    second = write_knowledge_document(files_ws, source, 'new knowledge', title='说明', file_id=first.file_id)
    assert second.file_id != first.file_id
    assert second.logical_type == 'knowledge_normalized'
    assert read_file_content(files_ws, first.file_id) == 'old knowledge'
    assert read_file_content(files_ws, second.file_id) == 'new knowledge'
    assert get_file_record(files_ws, first.file_id)['metadata']['historical_version']
    third = write_knowledge_document(files_ws, source, 'new knowledge', title='说明', file_id=second.file_id)
    assert third.file_id == second.file_id


def test_archive_member_conflict_is_rejected_before_creating_directory(files_ws):
    import zipfile
    from storage.file_store import import_user_upload
    from storage.source_workspace import extract_archive, move_source
    from storage.paths import workspace_root
    data = io.BytesIO()
    with zipfile.ZipFile(data, 'w') as archive:
        archive.writestr('config', 'a')
        archive.writestr('CONFIG/value.txt', 'b')
    record = import_user_upload(files_ws, io.BytesIO(data.getvalue()), 'conflict.zip')
    with pytest.raises(ValueError, match='file_directory_conflict'):
        extract_archive(files_ws, record.file_id, 'files/data/unpacked')
    assert not (workspace_root(files_ws) / 'files/data/unpacked').exists()
    with pytest.raises(ValueError, match='namespace_move'):
        move_source(files_ws, 'files/data', 'inbox/moved')


def test_backup_stream_and_reference_merge_preserve_single_owner_link(files_ws):
    from storage.file_bundle import export_bundle, restore_bundle
    from storage.file_store import import_user_upload
    from storage.reference_index import add_reference, list_references
    from storage.paths import workspace_root
    record = import_user_upload(files_ws, io.BytesIO(b'large stream payload'), 'stream.txt')
    add_reference(files_ws, record.file_id, 'custom', 'owner', 'source')
    output = io.BytesIO(); export_bundle(files_ws, output=output)
    (workspace_root(files_ws) / 'index/references.jsonl').write_text('')
    newer = add_reference(files_ws, record.file_id, 'custom', 'owner', 'source')
    assert restore_bundle(files_ws, io.BytesIO(output.getvalue()), preview=False)['ok']
    assert len(list_references(files_ws)) == 1 and list_references(files_ws)[0]['ref_id'] == newer.ref_id


def test_restore_unknown_only_readbacks_existing_facts(files_ws, monkeypatch):
    from storage.file_bundle import export_bundle, restore_bundle
    from storage.file_store import import_user_upload
    from storage.file_audit import reconcile_file_commits
    from storage.paths import workspace_root
    from storage import file_bundle
    record = import_user_upload(files_ws, io.BytesIO(b'original restore bytes'), 'restore.txt')
    bundle = export_bundle(files_ws)
    root = workspace_root(files_ws)
    (root / 'index/files.jsonl').write_text('')
    (root / record.path).unlink()
    from storage import records
    mutate = records.mutate_jsonl
    monkeypatch.setattr(records, 'mutate_jsonl', lambda *_: (_ for _ in ()).throw(RuntimeError('after restored payload')))
    with pytest.raises(RuntimeError):
        restore_bundle(files_ws, bundle, preview=False)
    monkeypatch.setattr(records, 'mutate_jsonl', mutate)
    assert (root / record.path).read_bytes() == b'original restore bytes'
    report = reconcile_file_commits(files_ws, apply=True)
    assert report['payload_writes'] == 0 and report['results'][0]['state'] == 'unresolved'
    assert not (root / 'index/files.jsonl').read_text()


def test_mutation_index_fault_keeps_actual_history_and_never_replays_writer(files_ws, monkeypatch):
    from storage.file_store import import_user_upload, resolve_file_path, read_file_content
    from storage.reference_index import add_reference
    from storage.file_mutations import managed_file_mutation, FileSettlementError
    from storage.file_audit import reconcile_file_commits
    from storage.project_changes import workspace_files_lock
    from storage import index
    record = import_user_upload(files_ws, io.BytesIO(b'old referenced bytes'), 'working.txt')
    add_reference(files_ws, record.file_id, 'message', 'session/run:user', 'attachment')
    path = resolve_file_path(files_ws, record.file_id)
    update = index.update_file_record
    monkeypatch.setattr(index, 'update_file_record', lambda *_: (_ for _ in ()).throw(RuntimeError('fault sealing history')))
    writes = 0
    with pytest.raises(FileSettlementError) as failure:
        with workspace_files_lock(files_ws), managed_file_mutation(files_ws, paths=[path]):
            writes += 1; path.write_bytes(b'new writer bytes')
    assert failure.value.as_result()['error_code'] == 'EXECUTION_UNKNOWN'
    monkeypatch.setattr(index, 'update_file_record', update)
    preview = reconcile_file_commits(files_ws)
    assert preview['results'][0]['writer_outcome'] == 'not_inferred'
    result = reconcile_file_commits(files_ws, apply=True)
    assert result['payload_writes'] == 0 and writes == 1
    assert read_file_content(files_ws, record.file_id) == 'old referenced bytes'
    current = result['results'][0]['file_changes'][0]['file_id']
    assert read_file_content(files_ws, current) == 'new writer bytes'
    assert not reconcile_file_commits(files_ws, apply=True)['results']


def test_unrecorded_external_overwrite_cannot_relabel_referenced_evidence(files_ws):
    from storage.file_store import import_user_upload, resolve_file_path, get_file_record
    from storage.reference_index import add_reference
    from storage.file_mutations import managed_file_mutation
    record = import_user_upload(files_ws, io.BytesIO(b'known evidence'), 'working.txt')
    add_reference(files_ws, record.file_id, 'message', 'session/run:user', 'attachment')
    path = resolve_file_path(files_ws, record.file_id); path.write_bytes(b'unrecorded overwrite')
    with pytest.raises(ValueError, match='referenced_payload_changed'):
        with managed_file_mutation(files_ws, paths=[path]):
            pytest.fail('writer must not run before original evidence is reconciled')
    assert get_file_record(files_ws, record.file_id)['sha256'] == record.sha256


@pytest.mark.parametrize('filename', ['broken.xlsx', 'broken.docx', 'broken.pptx', 'broken.pdf', 'broken.zip'])
def test_malformed_documents_report_processor_failure_and_preserve_original(files_ws, filename):
    from storage.file_store import import_user_upload, resolve_file_path
    from storage.file_formats import inspect_file
    record = import_user_upload(files_ws, io.BytesIO(b'not a valid document'), filename)
    with pytest.raises(ValueError):
        inspect_file(files_ws, record.file_id)
    assert resolve_file_path(files_ws, record.file_id).read_bytes() == b'not a valid document'
