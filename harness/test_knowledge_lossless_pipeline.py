"""Regression coverage for lossless knowledge ingestion and consumption."""

from __future__ import annotations

import io
import re
import zipfile


def _reset_context_runtime(tmp_path, monkeypatch):
    monkeypatch.setenv("LZCORE_WORKSPACE_ROOT", str(tmp_path / "workspaces"))
    monkeypatch.setenv("LZCORE_WORKSPACE_DIR", str(tmp_path / "workspaces"))
    import core.context.context_store as context_store
    import core.context.unified_retriever as unified_retriever

    context_store._stores.clear()
    unified_retriever._retrievers.clear()


def test_child_chunking_never_breaks_a_fenced_block():
    from agent.modules.knowledge.chunking import chunk_document
    from agent.modules.knowledge.schemas import NormalizedDocument

    code = "```\n" + ("display interface brief\n" * 160) + "```"
    doc = NormalizedDocument(
        source_id="ksrc_lossless", title="Runbook",
        normalized_markdown="# Runbook\n\nBefore.\n\n" + code + "\n\nAfter.",
    )
    parents, children = chunk_document(doc)

    assert "\n".join(parent.content for parent in parents).count(code) == 1
    assert all(child.content.count("```") in {0, 2} for child in children)
    assert sum(code in child.content for child in children) == 1


def test_html_and_docx_preserve_block_order_without_duplicate_text():
    from docx import Document
    from agent.modules.knowledge.parsers import docx, html

    parsed_html = html.parse(b"<blockquote><p>ONE_WARNING</p></blockquote>")
    assert parsed_html.normalized_markdown.count("ONE_WARNING") == 1

    document = Document()
    document.add_paragraph("BEFORE_TABLE")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE_CELL"
    document.add_paragraph("AFTER_TABLE")
    output = io.BytesIO()
    document.save(output)
    parsed_docx = docx.parse(output.getvalue()).normalized_markdown
    assert parsed_docx.index("BEFORE_TABLE") < parsed_docx.index("TABLE_CELL") < parsed_docx.index("AFTER_TABLE")


def test_xlsx_without_dimension_is_imported_without_row_loss():
    from openpyxl import Workbook
    from agent.modules.knowledge.parsers import xlsx

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["header"])
    sheet.append(["value"])
    raw = io.BytesIO()
    workbook.save(raw)
    repaired = io.BytesIO()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(repaired, "w") as target:
        for filename in source.namelist():
            body = source.read(filename)
            if filename == "xl/worksheets/sheet1.xml":
                body = re.sub(rb"<dimension\b[^>]*/>", b"", body)
            target.writestr(filename, body)

    parsed = xlsx.parse(repaired.getvalue())
    assert "value" in parsed.normalized_markdown
    assert parsed.warnings == []


def test_tool_reads_full_source_and_carries_retrieval_provenance(tmp_path, monkeypatch):
    _reset_context_runtime(tmp_path, monkeypatch)
    from agent.modules.knowledge.store import import_document
    from core.tools.general_tools.runtime_tools import (
        handle_knowledge_get_source,
        handle_knowledge_search,
    )
    from core.tools.schemas import ToolInvocation

    marker = "UNIQUE_LOSSLESS_MARKER"
    content = "# Global OSPF Guide\n\n" + marker + "\n\n" + ("route detail\n" * 500)
    imported = import_document(
        "knowledge_lossless_ws", "OSPF 配置指南", content,
        metadata={"scope": "global"},
    )
    assert imported["ok"] is True

    source_result = handle_knowledge_get_source(ToolInvocation(
        tool_id="knowledge.manage", workspace_id="knowledge_lossless_ws",
        arguments={"action": "read", "source_id": imported["source_id"], "level": "source"},
    ))
    assert source_result["content"] == content
    assert source_result["safe_excerpt"] == content
    assert source_result["scope"] == "global"
    assert source_result["chunk_count"] > 0

    search_result = handle_knowledge_search(ToolInvocation(
        tool_id="knowledge.manage", workspace_id="knowledge_lossless_ws",
        arguments={"action": "search", "query": marker, "scope": "global"},
    ))
    hit = search_result["results"][0]
    assert hit["source_id"] == imported["source_id"]
    assert hit["parent_chunk_id"]
    assert hit["scope"] == "global"
    assert marker in hit["safe_excerpt"]

    from agent.runtime.ssot_runtime import _build_retrieved_context_block
    context = _build_retrieved_context_block(
        workspace_id="knowledge_lossless_ws", session_id="session-1", task_id="task-1",
        user_input=marker, include_workspace_memory=False,
    )
    assert f"source_id={imported['source_id']}" in context
    assert "scope=global" in context
    assert "chunk_id=" in context and "parent_chunk_id=" in context


def test_title_fallback_preserves_scope_and_accepts_spaced_query(tmp_path, monkeypatch):
    _reset_context_runtime(tmp_path, monkeypatch)
    from agent.modules.knowledge.store import import_document
    from backend.api.knowledge_routes import _module_title_search

    imported = import_document(
        "knowledge_title_ws", "OSPF 配置指南", "unrelated body",
        metadata={"scope": "global"},
    )
    hits = _module_title_search(
        "knowledge_title_ws", "OSPF 配置指南", scope="global",
    )
    assert hits and hits[0]["source_id"] == imported["source_id"]
    assert hits[0]["scope"] == "global"
    assert _module_title_search("knowledge_title_ws", "OSPF 配置指南", scope="workspace") == []


def test_scope_filter_recovers_legacy_chunk_scope_from_its_source(tmp_path, monkeypatch):
    _reset_context_runtime(tmp_path, monkeypatch)
    from agent.modules.knowledge.store import import_document
    from core.context.context_store import get_context_store
    from core.context.unified_retriever import get_retriever

    imported = import_document(
        "knowledge_legacy_scope_ws", "Global Guide", "LEGACY_SCOPE_MARKER",
        metadata={"scope": "global"},
    )
    store = get_context_store("knowledge_legacy_scope_ws")
    for chunk in store.list_items(
        item_type="knowledge_chunk", source_id=imported["source_id"], limit=999,
    ):
        chunk["scope"] = "workspace"
        chunk["metadata"] = {
            key: value for key, value in (chunk.get("metadata") or {}).items()
            if key != "scope"
        }
        store.put(chunk)

    hits = get_retriever("knowledge_legacy_scope_ws").search_knowledge(
        "LEGACY_SCOPE_MARKER", scope="global",
    )
    assert hits and hits[0]["scope"] == "global"

    from agent.modules.knowledge.service import reindex_source
    assert reindex_source("knowledge_legacy_scope_ws", imported["source_id"])["ok"] is True
    assert get_retriever("knowledge_legacy_scope_ws").search_knowledge(
        "LEGACY_SCOPE_MARKER", scope="global",
    )
